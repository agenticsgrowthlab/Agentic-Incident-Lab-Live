from __future__ import annotations

import hashlib
import io
import math
import os
import uuid
from pathlib import Path
from typing import Any

from .core import client

KNOWLEDGE_DIR = Path(__file__).resolve().parent.parent / "knowledge"
EMBEDDING_MODEL = os.getenv("OPENAI_EMBEDDING_MODEL", "text-embedding-3-small")
EMBEDDING_DIMENSIONS = int(os.getenv("OPENAI_EMBEDDING_DIMENSIONS", "1536"))


def _documents() -> list[dict[str, str]]:
    documents = []
    for path in sorted(KNOWLEDGE_DIR.glob("*.md")):
        text = path.read_text(encoding="utf-8")
        lines = text.splitlines()
        title = lines[0].removeprefix("# ").strip() if lines else path.stem
        doc_type = "DOCUMENT"
        if len(lines) > 1 and lines[1].startswith("Type:"):
            doc_type = lines[1].split(":", 1)[1].strip().upper()
        documents.append(
            {"id": path.stem, "title": title, "type": doc_type, "content": text}
        )
    return documents


def _embed(texts: list[str]) -> list[list[float]]:
    response = client().embeddings.create(
        model=EMBEDDING_MODEL,
        input=texts,
        dimensions=EMBEDDING_DIMENSIONS,
    )
    return [item.embedding for item in response.data]


def _cosine(left: list[float], right: list[float]) -> float:
    dot = sum(a * b for a, b in zip(left, right))
    left_norm = math.sqrt(sum(value * value for value in left))
    right_norm = math.sqrt(sum(value * value for value in right))
    return dot / (left_norm * right_norm) if left_norm and right_norm else 0.0


def _memory_search(
    query: str, documents: list[dict[str, str]]
) -> list[dict[str, Any]]:
    vectors = _embed([query, *[doc["content"] for doc in documents]])
    query_vector, doc_vectors = vectors[0], vectors[1:]
    ranked = sorted(
        [
            dict(doc, score=max(0.0, _cosine(query_vector, vector)))
            for doc, vector in zip(documents, doc_vectors)
        ],
        key=lambda item: item["score"],
        reverse=True,
    )
    return ranked[:6]


def _vector_literal(vector: list[float]) -> str:
    return "[" + ",".join(f"{value:.8f}" for value in vector) + "]"


def _ensure_vector_tables(cursor: Any) -> None:
    cursor.execute("CREATE EXTENSION IF NOT EXISTS vector")
    cursor.execute(
        f"""
        CREATE TABLE IF NOT EXISTS incident_knowledge (
            id text PRIMARY KEY,
            title text NOT NULL,
            doc_type text NOT NULL,
            content text NOT NULL,
            content_hash text NOT NULL,
            embedding vector({EMBEDDING_DIMENSIONS}) NOT NULL
        )
        """
    )
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS uploaded_knowledge_documents (
            id text PRIMARY KEY,
            title text NOT NULL,
            filename text NOT NULL,
            content_type text NOT NULL,
            doc_type text NOT NULL,
            incident_id text,
            content_hash text NOT NULL,
            extracted_text text NOT NULL,
            original_bytes bytea NOT NULL,
            chunk_count integer NOT NULL DEFAULT 0,
            created_at timestamptz NOT NULL DEFAULT now()
        )
        """
    )
    cursor.execute(
        f"""
        CREATE TABLE IF NOT EXISTS uploaded_knowledge_chunks (
            id text PRIMARY KEY,
            document_id text NOT NULL
              REFERENCES uploaded_knowledge_documents(id) ON DELETE CASCADE,
            chunk_index integer NOT NULL,
            content text NOT NULL,
            embedding vector({EMBEDDING_DIMENSIONS}) NOT NULL,
            created_at timestamptz NOT NULL DEFAULT now(),
            UNIQUE(document_id, chunk_index)
        )
        """
    )


def _sync_bundled_documents(cursor: Any, documents: list[dict[str, str]]) -> None:
    cursor.execute("SELECT id, content_hash FROM incident_knowledge")
    existing = dict(cursor.fetchall())
    changed: list[tuple[dict[str, str], str]] = []

    for doc in documents:
        digest = hashlib.sha256(doc["content"].encode("utf-8")).hexdigest()
        if existing.get(doc["id"]) != digest:
            changed.append((doc, digest))

    if not changed:
        return

    embeddings = _embed([doc["content"] for doc, _ in changed])
    for (doc, digest), embedding in zip(changed, embeddings):
        cursor.execute(
            """
            INSERT INTO incident_knowledge
              (id, title, doc_type, content, content_hash, embedding)
            VALUES (%s, %s, %s, %s, %s, %s::vector)
            ON CONFLICT (id) DO UPDATE SET
              title=EXCLUDED.title,
              doc_type=EXCLUDED.doc_type,
              content=EXCLUDED.content,
              content_hash=EXCLUDED.content_hash,
              embedding=EXCLUDED.embedding
            """,
            (
                doc["id"],
                doc["title"],
                doc["type"],
                doc["content"],
                digest,
                _vector_literal(embedding),
            ),
        )


def _pgvector_search(
    query: str, documents: list[dict[str, str]]
) -> list[dict[str, Any]]:
    import psycopg

    database_url = os.environ["DATABASE_URL"]
    with psycopg.connect(database_url, autocommit=True) as connection:
        with connection.cursor() as cursor:
            _ensure_vector_tables(cursor)
            _sync_bundled_documents(cursor, documents)

            query_embedding = _embed([query])[0]
            literal = _vector_literal(query_embedding)

            cursor.execute(
                """
                WITH candidates AS (
                    SELECT
                        id,
                        title,
                        doc_type,
                        content,
                        1 - (embedding <=> %s::vector) AS score
                    FROM incident_knowledge

                    UNION ALL

                    SELECT
                        c.id,
                        d.title,
                        d.doc_type,
                        c.content,
                        1 - (c.embedding <=> %s::vector) AS score
                    FROM uploaded_knowledge_chunks c
                    JOIN uploaded_knowledge_documents d
                      ON d.id = c.document_id
                )
                SELECT id, title, doc_type, content, score
                FROM candidates
                ORDER BY score DESC
                LIMIT 6
                """,
                (literal, literal),
            )

            return [
                {
                    "id": row[0],
                    "title": row[1],
                    "type": row[2],
                    "content": row[3],
                    "score": max(0.0, float(row[4])),
                }
                for row in cursor.fetchall()
            ]


def _decode_text(raw: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("The file could not be decoded as text")


def _extract_text(filename: str, raw: bytes) -> str:
    suffix = Path(filename).suffix.lower()

    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw))
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == ".docx":
        from docx import Document

        document = Document(io.BytesIO(raw))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                text += "\n" + " | ".join(cell.text for cell in row.cells)
    elif suffix in {".txt", ".md", ".csv", ".json", ".yaml", ".yml"}:
        text = _decode_text(raw)
    else:
        raise ValueError(
            "Unsupported file type. Use PDF, DOCX, TXT, Markdown, CSV, JSON, or YAML."
        )

    cleaned = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if len(cleaned) < 20:
        raise ValueError("The uploaded document did not contain enough extractable text")
    return cleaned


def _chunk_text(text: str, target: int = 1400, overlap: int = 200) -> list[str]:
    paragraphs = [part.strip() for part in text.split("\n\n") if part.strip()]
    if not paragraphs:
        paragraphs = [text]

    chunks: list[str] = []
    current = ""

    for paragraph in paragraphs:
        candidate = f"{current}\n\n{paragraph}".strip() if current else paragraph
        if len(candidate) <= target:
            current = candidate
            continue

        if current:
            chunks.append(current)

        if len(paragraph) <= target:
            current = paragraph
            continue

        start = 0
        while start < len(paragraph):
            end = min(len(paragraph), start + target)
            chunks.append(paragraph[start:end].strip())
            if end == len(paragraph):
                current = ""
                break
            start = max(0, end - overlap)

    if current:
        chunks.append(current)

    return [chunk for chunk in chunks if chunk]


def ingest_uploaded_document(
    *,
    filename: str,
    content_type: str,
    raw: bytes,
    doc_type: str = "INCIDENT_EVIDENCE",
    incident_id: str | None = None,
) -> dict[str, Any]:
    if not os.getenv("DATABASE_URL"):
        raise RuntimeError("DATABASE_URL is not configured")

    import psycopg

    text = _extract_text(filename, raw)
    digest = hashlib.sha256(raw).hexdigest()
    title = (
        Path(filename).stem.replace("_", " ").replace("-", " ").strip()
        or filename
    )
    chunks = _chunk_text(text)

    if len(chunks) > 500:
        raise ValueError("The document is too large after chunking")

    with psycopg.connect(os.environ["DATABASE_URL"], autocommit=True) as connection:
        with connection.cursor() as cursor:
            _ensure_vector_tables(cursor)

            cursor.execute(
                """
                SELECT id, title, filename, content_type, doc_type, incident_id,
                       chunk_count, created_at
                FROM uploaded_knowledge_documents
                WHERE content_hash=%s
                ORDER BY created_at DESC
                LIMIT 1
                """,
                (digest,),
            )
            duplicate = cursor.fetchone()
            if duplicate:
                return {
                    "id": duplicate[0],
                    "title": duplicate[1],
                    "filename": duplicate[2],
                    "content_type": duplicate[3],
                    "doc_type": duplicate[4],
                    "incident_id": duplicate[5],
                    "chunk_count": duplicate[6],
                    "created_at": duplicate[7].isoformat(),
                    "status": "already_indexed",
                }

            document_id = f"doc_{uuid.uuid4().hex[:16]}"
            embeddings = _embed(chunks)

            cursor.execute(
                """
                INSERT INTO uploaded_knowledge_documents
                  (id, title, filename, content_type, doc_type, incident_id,
                   content_hash, extracted_text, original_bytes, chunk_count)
                VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                """,
                (
                    document_id,
                    title,
                    filename,
                    content_type,
                    doc_type.upper(),
                    incident_id,
                    digest,
                    text,
                    raw,
                    len(chunks),
                ),
            )

            for index, (chunk, embedding) in enumerate(zip(chunks, embeddings)):
                cursor.execute(
                    """
                    INSERT INTO uploaded_knowledge_chunks
                      (id, document_id, chunk_index, content, embedding)
                    VALUES (%s,%s,%s,%s,%s::vector)
                    """,
                    (
                        f"{document_id}_chunk_{index}",
                        document_id,
                        index,
                        chunk,
                        _vector_literal(embedding),
                    ),
                )

            cursor.execute(
                "SELECT created_at FROM uploaded_knowledge_documents WHERE id=%s",
                (document_id,),
            )
            created_at = cursor.fetchone()[0]

    return {
        "id": document_id,
        "title": title,
        "filename": filename,
        "content_type": content_type,
        "doc_type": doc_type.upper(),
        "incident_id": incident_id,
        "chunk_count": len(chunks),
        "created_at": created_at.isoformat(),
        "status": "indexed",
    }


def list_uploaded_documents() -> list[dict[str, Any]]:
    if not os.getenv("DATABASE_URL"):
        raise RuntimeError("DATABASE_URL is not configured")

    import psycopg

    with psycopg.connect(os.environ["DATABASE_URL"], autocommit=True) as connection:
        with connection.cursor() as cursor:
            _ensure_vector_tables(cursor)
            cursor.execute(
                """
                SELECT id, title, filename, content_type, doc_type, incident_id,
                       chunk_count, created_at
                FROM uploaded_knowledge_documents
                ORDER BY created_at DESC
                LIMIT 250
                """
            )
            return [
                {
                    "id": row[0],
                    "title": row[1],
                    "filename": row[2],
                    "content_type": row[3],
                    "doc_type": row[4],
                    "incident_id": row[5],
                    "chunk_count": row[6],
                    "created_at": row[7].isoformat(),
                    "status": "indexed",
                }
                for row in cursor.fetchall()
            ]


def retrieve(query: str) -> tuple[list[dict[str, Any]], str, str | None]:
    documents = _documents()

    if os.getenv("DATABASE_URL"):
        try:
            return (
                _pgvector_search(query, documents),
                "Neon PostgreSQL + pgvector · bundled + uploaded knowledge",
                None,
            )
        except Exception as exc:
            if not documents:
                raise
            evidence = _memory_search(query, documents)
            return (
                evidence,
                "OpenAI embeddings · bundled in-memory fallback",
                f"pgvector fallback: {type(exc).__name__}",
            )

    if not documents:
        raise RuntimeError("No knowledge documents are available")

    return (
        _memory_search(query, documents),
        "OpenAI embeddings · bundled in-memory vectors",
        None,
    )
